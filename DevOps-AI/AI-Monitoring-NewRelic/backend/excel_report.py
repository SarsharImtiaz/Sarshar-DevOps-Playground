"""Builds the formatted Excel workbook (dashboard + detail sheets) from an Analysis."""
from io import BytesIO

import pandas as pd
from openpyxl import Workbook
from openpyxl.cell.cell import ILLEGAL_CHARACTERS_RE
from openpyxl.chart import BarChart, LineChart, Reference
from openpyxl.chart.label import DataLabelList
from openpyxl.chart.shapes import GraphicalProperties
from openpyxl.drawing.line import LineProperties
from openpyxl.drawing.text import CharacterProperties, ParagraphProperties
from openpyxl.formatting.rule import ColorScaleRule, DataBarRule, FormulaRule
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.table import Table, TableStyleInfo

from analysis import CATEGORY_DESCRIPTIONS

# ---- Theme ---------------------------------------------------------------
NAVY = "1F3864"
INK = "0B0B0B"
MUTED = "52514E"
CARD_FILL = "F3F6FA"
RULE = "D9DEE7"
ACCENT = "2A78D6"

# Fixed colour per category - colour follows the category on every chart
CATEGORY_COLORS = {
    "HTTP 5xx Response": "2A78D6",
    "Request Context Log": "1BAF7A",
    "Authentication": "EDA100",
    "Database": "E87BA4",
    "Unhandled Exception": "008300",
    "Application Error": "E34948",
    "Validation": "4A3AA7",
    "HTTP 4xx Response": "EB6834",
}

# Sequential blue ramp (light -> dark) for heatmaps
HEAT_LOW, HEAT_MID, HEAT_HIGH = "FFFFFF", "9EC5F4", "3987E5"

TABLE_STYLE = "TableStyleMedium2"
FMT_INT = "#,##0"
FMT_PCT = "0.0%"
FMT_DT = "yyyy-mm-dd hh:mm"
FMT_DATE = "ddd dd mmm yyyy"

TITLE_FONT = Font(name="Calibri", size=18, bold=True, color=NAVY)
SUBTITLE_FONT = Font(name="Calibri", size=10, italic=True, color=MUTED)
SECTION_FONT = Font(name="Calibri", size=13, bold=True, color=NAVY)
HEADER_FONT = Font(name="Calibri", size=10, bold=True, color="FFFFFF")
HEADER_FILL = PatternFill("solid", fgColor=NAVY)
THIN = Side(style="thin", color=RULE)
BOX = Border(left=THIN, right=THIN, top=THIN, bottom=THIN)
WRAP_TOP = Alignment(wrap_text=True, vertical="top")
TOP = Alignment(vertical="top")

MAX_CELL = 32000


# ---- Cell helpers --------------------------------------------------------
def _value(v):
    if v is None:
        return None
    if isinstance(v, float) and pd.isna(v):
        return None
    if v is pd.NaT:
        return None
    if isinstance(v, pd.Timestamp):
        return v.to_pydatetime()
    if isinstance(v, str):
        return ILLEGAL_CHARACTERS_RE.sub("", v)[:MAX_CELL]
    if hasattr(v, "item"):  # numpy scalar
        return v.item()
    return v


def _set(ws, row, col, value):
    cell = ws.cell(row=row, column=col, value=_value(value))
    if isinstance(cell.value, str) and cell.value.startswith("="):
        cell.data_type = "s"  # never let log text be read as a formula
    return cell


def _title(ws, title, subtitle, width_cols=10):
    ws["A1"] = title
    ws["A1"].font = TITLE_FONT
    ws["A2"] = subtitle
    ws["A2"].font = SUBTITLE_FONT
    ws.row_dimensions[1].height = 28
    ws.sheet_view.showGridLines = False


def write_table(ws, df, name, top=4, widths=None, formats=None, wrap=(), databars=(), freeze=True,
                set_widths=True):
    """Writes df as a styled, filterable Excel Table with its header on row `top`."""
    widths, formats = widths or {}, formats or {}
    headers = [str(c) for c in df.columns]
    for c, h in enumerate(headers, start=1):
        _set(ws, top, c, h)

    rows = df.itertuples(index=False, name=None)
    for r, row in enumerate(rows, start=top + 1):
        for c, v in enumerate(row, start=1):
            _set(ws, r, c, v)

    last = top + max(len(df), 1)
    ref = f"A{top}:{get_column_letter(len(headers))}{last}"
    table = Table(displayName=name, ref=ref)
    table.tableStyleInfo = TableStyleInfo(name=TABLE_STYLE, showRowStripes=True)
    ws.add_table(table)

    ws.row_dimensions[top].height = 30
    for c, h in enumerate(headers, start=1):
        letter = get_column_letter(c)
        ws.cell(top, c).alignment = Alignment(wrap_text=True, vertical="center")
        if set_widths:  # False for a second table stacked under the first on the same sheet
            ws.column_dimensions[letter].width = widths.get(h, max(10, min(len(h) + 4, 40)))
        fmt = formats.get(h)
        align = WRAP_TOP if h in wrap else TOP
        for r in range(top + 1, last + 1):
            cell = ws.cell(r, c)
            cell.alignment = align
            if fmt:
                cell.number_format = fmt

    for h in databars:
        if h in headers and len(df):
            letter = get_column_letter(headers.index(h) + 1)
            ws.conditional_formatting.add(
                f"{letter}{top + 1}:{letter}{last}",
                DataBarRule(start_type="num", start_value=0, end_type="max", color=ACCENT),
            )

    if freeze:
        ws.freeze_panes = ws.cell(top + 1, 1)
    return top, last


def _col(df, name):
    return list(df.columns).index(name) + 1


# ---- Charts --------------------------------------------------------------
def _set_title(chart, text):
    """Chart title that reserves its own space instead of overlaying the plot."""
    chart.title = text
    chart.title.overlay = False
    chart.title.tx.rich.p[0].pPr = ParagraphProperties(defRPr=CharacterProperties(sz=1200, b=True, solidFill=NAVY))


def _style_axes(chart):
    chart.x_axis.delete = False
    chart.y_axis.delete = False
    chart.y_axis.majorGridlines.spPr = GraphicalProperties(ln=LineProperties(solidFill="E7E9EE"))
    chart.x_axis.spPr = GraphicalProperties(ln=LineProperties(solidFill="BFC4CC"))
    chart.y_axis.spPr = GraphicalProperties(ln=LineProperties(noFill=True))


def _fill(series, color):
    series.graphicalProperties.solidFill = color
    series.graphicalProperties.line.solidFill = color


def hbar_chart(title, data_ws, label_col, value_col, first_row, last_row, color=ACCENT,
               colors=None, width=18, height=None):
    """Horizontal bar chart, largest at the top, values labelled at bar ends."""
    chart = BarChart()
    chart.type = "bar"
    chart.style = 10
    _set_title(chart, title)
    chart.legend = None
    chart.gapWidth = 60
    data = Reference(data_ws, min_col=value_col, min_row=first_row - 1, max_row=last_row)
    cats = Reference(data_ws, min_col=label_col, min_row=first_row, max_row=last_row)
    chart.add_data(data, titles_from_data=True)
    chart.set_categories(cats)
    chart.x_axis.scaling.orientation = "maxMin"  # rank 1 at the top
    chart.y_axis.crosses = "max"  # keep the value axis at the bottom when categories are reversed
    chart.y_axis.number_format = FMT_INT
    _style_axes(chart)
    series = chart.series[0]
    _fill(series, color)
    if colors:
        from openpyxl.chart.marker import DataPoint
        for idx, c in enumerate(colors):
            pt = DataPoint(idx=idx)
            pt.graphicalProperties.solidFill = c
            pt.graphicalProperties.line.solidFill = c
            series.dPt.append(pt)
    series.dLbls = DataLabelList()
    series.dLbls.showVal = True
    for flag in ("showSerName", "showCatName", "showLegendKey", "showPercent"):
        setattr(series.dLbls, flag, False)
    series.dLbls.numFmt = FMT_INT
    series.dLbls.position = "outEnd"
    chart.width = width
    chart.height = height or max(6, 0.55 * (last_row - first_row + 1) + 2)
    return chart


# ---- Sheets --------------------------------------------------------------
def _overview(wb, a, data):
    ws = wb["Overview"]
    ws.sheet_view.showGridLines = False
    ws.sheet_properties.tabColor = NAVY
    ws.column_dimensions["A"].width = 2
    for c in range(2, 15):
        ws.column_dimensions[get_column_letter(c)].width = 12

    ws["B2"] = f"New Relic Error Report - {a.service}"
    ws["B2"].font = Font(name="Calibri", size=22, bold=True, color=NAVY)
    ws.row_dimensions[2].height = 34
    ws["B3"] = (f"Account {a.account_id}   |   Period (UTC): {a.period_start:%d %b %Y %H:%M} - "
                f"{a.period_end:%d %b %Y %H:%M}   |   Generated: {a.generated_at:%d %b %Y %H:%M} UTC   |   "
                f"Level: ERROR")
    ws["B3"].font = SUBTITLE_FONT

    # KPI cards: 4 per row, 3 columns wide each
    cards = list(a.kpis.items())
    for i, (label, value) in enumerate(cards):
        r = 5 + (i // 4) * 3
        c = 2 + (i % 4) * 3
        ws.merge_cells(start_row=r, start_column=c, end_row=r, end_column=c + 2)
        ws.merge_cells(start_row=r + 1, start_column=c, end_row=r + 1, end_column=c + 2)
        lab = ws.cell(r, c, label.upper())
        lab.font = Font(size=9, bold=True, color=MUTED)
        val = ws.cell(r + 1, c, _value(value))
        val.font = Font(size=18 if isinstance(value, (int, float)) else 13, bold=True, color=NAVY)
        if isinstance(value, (int, float)):
            val.number_format = FMT_INT
        for rr in (r, r + 1):
            for cc in range(c, c + 3):
                cell = ws.cell(rr, cc)
                cell.fill = PatternFill("solid", fgColor=CARD_FILL)
                cell.alignment = Alignment(horizontal="left", vertical="center", indent=1)
        for cc in range(c, c + 3):
            ws.cell(r, cc).border = Border(top=Side(style="medium", color=ACCENT))
        ws.row_dimensions[r + 1].height = 30

    # Key findings
    r = 5 + ((len(cards) + 3) // 4) * 3 + 1
    ws.cell(r, 2, "Key Findings").font = SECTION_FONT
    r += 1
    for c, h, span in ((2, "Area", 2), (4, "Finding", 10)):
        ws.merge_cells(start_row=r, start_column=c, end_row=r, end_column=c + span - 1)
        cell = ws.cell(r, c, h)
        for cc in range(c, c + span):
            ws.cell(r, cc).fill = HEADER_FILL
        cell.font = HEADER_FONT
    for i, (area, text) in enumerate(a.findings):
        r += 1
        ws.merge_cells(start_row=r, start_column=2, end_row=r, end_column=3)
        ws.merge_cells(start_row=r, start_column=4, end_row=r, end_column=13)
        ws.cell(r, 2, area).font = Font(bold=True, color=INK)
        _set(ws, r, 4, text)
        ws.cell(r, 4).alignment = WRAP_TOP
        ws.cell(r, 2).alignment = TOP
        ws.row_dimensions[r].height = 15 * max(1, -(-len(text) // 125))
        if i % 2:
            for cc in range(2, 14):
                ws.cell(r, cc).fill = PatternFill("solid", fgColor=CARD_FILL)
        for cc in range(2, 14):
            ws.cell(r, cc).border = Border(bottom=THIN)

    # Errors by category (table left, chart below)
    r += 2
    ws.cell(r, 2, "Errors by Category").font = SECTION_FONT
    r += 1
    heads = [("Category", 3), ("Errors", 1), ("Share", 1), ("Patterns", 1), ("Tenants", 1), ("What it means", 6)]
    c = 2
    for h, span in heads:
        if span > 1:
            ws.merge_cells(start_row=r, start_column=c, end_row=r, end_column=c + span - 1)
        ws.cell(r, c, h).font = HEADER_FONT
        for cc in range(c, c + span):
            ws.cell(r, cc).fill = HEADER_FILL
        c += span
    for _, row in a.categories.iterrows():
        r += 1
        ws.merge_cells(start_row=r, start_column=2, end_row=r, end_column=4)
        ws.merge_cells(start_row=r, start_column=9, end_row=r, end_column=14)
        ws.cell(r, 2, row["Category"]).font = Font(bold=True, color=INK)
        ws.cell(r, 5, int(row["Errors"])).number_format = FMT_INT
        ws.cell(r, 6, float(row["Share"])).number_format = FMT_PCT
        ws.cell(r, 7, int(row["Patterns"])).number_format = FMT_INT
        ws.cell(r, 8, int(row["Tenants"])).number_format = FMT_INT
        ws.cell(r, 9, row["Description"])
        ws.row_dimensions[r].height = 30
        for cc in range(2, 15):
            ws.cell(r, cc).border = Border(bottom=THIN)
            ws.cell(r, cc).alignment = WRAP_TOP if cc == 9 else TOP
        # colour marker matching the category's colour on every chart
        ws.cell(r, 2).border = Border(left=Side(style="thick", color=CATEGORY_COLORS.get(row["Category"], ACCENT)),
                                      bottom=THIN)
        ws.cell(r, 2).alignment = Alignment(vertical="top", indent=1)
    cat_end = r

    # Charts
    r = cat_end + 2
    ws.cell(r, 2, "Daily Errors by Category").font = SECTION_FONT
    ws.add_chart(_daily_chart(wb, a), f"B{r + 1}")


def _daily_chart(wb, a):
    ws = wb["Daily Trend"]
    n = len(a.daily)
    chart = BarChart()
    chart.type = "col"
    chart.grouping = "stacked"
    chart.overlap = 100
    chart.gapWidth = 40
    _set_title(chart, "Errors per day (stacked by category)")
    categories = [c for c in a.daily.columns if c not in ("Date", "Total")]
    for name in categories:
        col = _col(a.daily, name)
        chart.add_data(Reference(ws, min_col=col, min_row=4, max_row=4 + n), titles_from_data=True)
        _fill(chart.series[-1], CATEGORY_COLORS.get(name, ACCENT))
    chart.set_categories(Reference(ws, min_col=1, min_row=5, max_row=4 + n))
    chart.x_axis.number_format = "ddd dd mmm"
    chart.y_axis.number_format = FMT_INT
    chart.legend.position = "t"
    _style_axes(chart)
    chart.width, chart.height = 30, 10
    return chart


def _chart_data(wb, a):
    """Hidden sheet holding short labels for chart axes."""
    ws = wb.create_sheet("ChartData")
    ws.sheet_state = "hidden"
    out = {}

    def block(col, headers, rows):
        for i, h in enumerate(headers):
            ws.cell(1, col + i, h)
        for r, row in enumerate(rows, start=2):
            for i, v in enumerate(row):
                cell = _set(ws, r, col + i, v)
                if isinstance(cell.value, (int, float)):
                    cell.number_format = FMT_INT
        return 2, max(2, len(rows) + 1)

    out["categories"] = block(1, ["Category", "Errors"],
                              list(zip(a.categories["Category"], a.categories["Errors"])))
    pats = a.patterns.head(15)
    out["patterns"] = block(4, ["Pattern", "Errors"],
                            [(f"{p.PatternId}  {_short(p.Pattern, 55)}", p.Errors) for p in pats.itertuples()])
    out["pattern_colors"] = [CATEGORY_COLORS.get(c, ACCENT) for c in pats["Category"]]
    ten = a.tenants.head(15)
    out["tenants"] = block(7, ["Tenant", "Errors"], list(zip(ten["TenantName"], ten["Errors"])))
    eps = a.endpoints.head(15)
    out["endpoints"] = block(10, ["Endpoint", "Errors"], list(zip(eps["Endpoint"], eps["Errors"])))
    out["hourly"] = block(13, ["Hour", "Errors"],
                          [(f"{h:%a %d %b}" if h.hour == 0 else f"{h:%d %b %H:00}", n)
                           for h, n in zip(a.hourly["Hour"], a.hourly["Errors"])])
    return out


def _short(text, n):
    return text if len(text) <= n else text[: n - 1] + "…"


def _charts_sheet(wb, a, data):
    ws = wb.create_sheet("Charts", 1)
    ws.sheet_properties.tabColor = NAVY
    _title(ws, "Charts", "Visual summary. Colours follow the category and match the Overview. Times are UTC.")
    cd = wb["ChartData"]

    # Hourly trend
    first, last = data["hourly"]
    line = LineChart()
    _set_title(line, "Errors per hour")
    line.add_data(Reference(cd, min_col=14, min_row=1, max_row=last), titles_from_data=True)
    line.set_categories(Reference(cd, min_col=13, min_row=2, max_row=last))
    line.legend = None
    s = line.series[0]
    s.graphicalProperties.line.solidFill = ACCENT
    s.graphicalProperties.line.width = 22000
    s.smooth = False
    # hourly data starts at midnight, so every 24th label is the start of a day
    line.x_axis.tickLblSkip = 24
    line.x_axis.tickMarkSkip = 24
    line.x_axis.number_format = "@"
    line.y_axis.number_format = FMT_INT
    _style_axes(line)
    line.width, line.height = 32, 9
    ws.add_chart(line, "A4")

    first, last = data["patterns"]
    ws.add_chart(hbar_chart("Top 15 error patterns", cd, 4, 5, first, last,
                            colors=data["pattern_colors"], width=32), "A23")
    r = 23 + int(max(6, 0.55 * (last - first + 1) + 2) * 2) + 2
    first, last = data["tenants"]
    ws.add_chart(hbar_chart("Top 15 tenants by errors", cd, 7, 8, first, last, width=16), f"A{r}")
    first, last = data["endpoints"]
    ws.add_chart(hbar_chart("Top 15 endpoints by errors", cd, 10, 11, first, last, width=16), f"K{r}")


def _patterns_sheet(wb, a):
    ws = wb.create_sheet("Error Patterns")
    ws.sheet_properties.tabColor = ACCENT
    _title(ws, "Error Patterns",
           "Each row is one error type. Similar messages are grouped: IDs, timestamps, e-mails and numbers are "
           "replaced by {id}, {time}, {email}, {n}. Use the Pattern ID to filter Raw Logs.")
    df = a.patterns[["PatternId", "Category", "Pattern", "Errors", "Share", "Tenants", "TopTenant",
                     "TopEndpoint", "FirstSeen", "LastSeen", "Sample"]].rename(columns={
        "PatternId": "Pattern ID", "Share": "% of Total", "Tenants": "Tenants Affected",
        "TopTenant": "Most Affected Tenant", "TopEndpoint": "Most Affected Endpoint",
        "FirstSeen": "First Seen (UTC)", "LastSeen": "Last Seen (UTC)", "Sample": "Sample Message",
    })
    df["Sample Message"] = df["Sample Message"].str.slice(0, 1500)
    write_table(ws, df, "ErrorPatterns",
                widths={"Pattern ID": 10, "Category": 20, "Pattern": 70, "Errors": 10, "% of Total": 10,
                        "Tenants Affected": 10, "Most Affected Tenant": 28, "Most Affected Endpoint": 32,
                        "First Seen (UTC)": 17, "Last Seen (UTC)": 17, "Sample Message": 80},
                formats={"Errors": FMT_INT, "% of Total": FMT_PCT, "First Seen (UTC)": FMT_DT,
                         "Last Seen (UTC)": FMT_DT},
                wrap=("Pattern", "Most Affected Tenant"), databars=("Errors",))
    _category_markers(ws, df, "Category", 4)


def _category_markers(ws, df, column, top):
    """Thick left border in the category colour on the Category column."""
    col = get_column_letter(_col(df, column))
    rng = f"{col}{top + 1}:{col}{top + max(len(df), 1)}"
    for cat, color in CATEGORY_COLORS.items():
        ws.conditional_formatting.add(rng, FormulaRule(
            formula=[f'${col}{top + 1}="{cat}"'],
            font=Font(color=INK, bold=True),
            border=Border(left=Side(style="thick", color=color))))


def _tenants_sheet(wb, a):
    ws = wb.create_sheet("Tenants")
    ws.sheet_properties.tabColor = "1BAF7A"
    _title(ws, "Tenant Impact",
           f"Errors per tenant. {a.stats['tenant_unattributed']:,} errors with no tenant information are "
           "excluded (see Data Quality). 'Attribution' shows how the tenant was identified.")
    df = a.tenants[["Rank", "TenantName", "TenantId", "Source", "Errors", "Share", "Patterns", "TopPattern",
                    "TopEndpoint", "FirstSeen", "LastSeen"]].rename(columns={
        "TenantName": "Tenant", "TenantId": "Tenant ID", "Source": "Attribution", "Share": "% of Total",
        "Patterns": "Distinct Patterns", "TopPattern": "Most Frequent Pattern", "TopEndpoint": "Most Affected Endpoint",
        "FirstSeen": "First Seen (UTC)", "LastSeen": "Last Seen (UTC)",
    })
    write_table(ws, df, "Tenants",
                widths={"Rank": 7, "Tenant": 34, "Tenant ID": 39, "Attribution": 24, "Errors": 10,
                        "% of Total": 10, "Distinct Patterns": 10, "Most Frequent Pattern": 60,
                        "Most Affected Endpoint": 32, "First Seen (UTC)": 17, "Last Seen (UTC)": 17},
                formats={"Errors": FMT_INT, "% of Total": FMT_PCT, "First Seen (UTC)": FMT_DT,
                         "Last Seen (UTC)": FMT_DT},
                wrap=("Most Frequent Pattern",), databars=("Errors",))


def _tenant_top_sheet(wb, a):
    ws = wb.create_sheet("Tenant Top Errors")
    ws.sheet_properties.tabColor = "1BAF7A"
    _title(ws, "Top 5 Error Patterns per Tenant",
           "For each tenant, its five most frequent error patterns. Tenants are ordered by total errors.")
    df = a.tenant_top_patterns[["TenantName", "Rank", "PatternId", "Category", "Pattern", "Errors",
                                "ShareOfTenant"]].rename(columns={
        "TenantName": "Tenant", "PatternId": "Pattern ID", "ShareOfTenant": "% of Tenant's Errors"})
    top, last = write_table(ws, df, "TenantTopErrors",
                            widths={"Tenant": 34, "Rank": 7, "Pattern ID": 10, "Category": 20, "Pattern": 80,
                                    "Errors": 10, "% of Tenant's Errors": 12},
                            formats={"Errors": FMT_INT, "% of Tenant's Errors": FMT_PCT},
                            wrap=("Pattern",), databars=("% of Tenant's Errors",))
    # Separator line + bold name at the start of each tenant group
    ws.conditional_formatting.add(f"A{top + 1}:G{last}", FormulaRule(
        formula=[f"$B{top + 1}=1"], font=Font(bold=True),
        border=Border(top=Side(style="medium", color=NAVY))))


def _matrix_sheet(wb, df, sheet, table, title, subtitle, first_col_width, tab):
    ws = wb.create_sheet(sheet)
    ws.sheet_properties.tabColor = tab
    _title(ws, title, subtitle)
    first = df.columns[0]
    widths = {c: 9 for c in df.columns}
    widths[first] = first_col_width
    widths["Total"] = 10
    formats = {c: FMT_INT for c in df.columns[1:]}
    if first == "Date":
        formats["Date"] = FMT_DATE
    top, last = write_table(ws, df, table, widths=widths, formats=formats)
    ws.freeze_panes = ws.cell(top + 1, 2)
    if len(df):
        rng = f"B{top + 1}:{get_column_letter(len(df.columns) - 1)}{last}"
        ws.conditional_formatting.add(rng, ColorScaleRule(
            start_type="num", start_value=0, start_color=HEAT_LOW,
            mid_type="percentile", mid_value=75, mid_color=HEAT_MID,
            end_type="max", end_color=HEAT_HIGH))
        tot = get_column_letter(len(df.columns))
        ws.conditional_formatting.add(f"{tot}{top + 1}:{tot}{last}", DataBarRule(
            start_type="num", start_value=0, end_type="max", color=ACCENT))
    return ws


def _endpoints_sheet(wb, a):
    ws = wb.create_sheet("Endpoints")
    ws.sheet_properties.tabColor = "EDA100"
    _title(ws, "Endpoint Impact",
           "Errors per API route. IDs in the path are replaced by {id} so all calls to the same route are grouped.")
    df = a.endpoints[["Rank", "Endpoint", "Errors", "Share", "Tenants", "Patterns", "TopPattern"]].rename(columns={
        "Share": "% of Total", "Tenants": "Tenants Affected", "Patterns": "Distinct Patterns",
        "TopPattern": "Most Frequent Pattern"})
    write_table(ws, df, "Endpoints",
                widths={"Rank": 7, "Endpoint": 45, "Errors": 10, "% of Total": 10, "Tenants Affected": 10,
                        "Distinct Patterns": 10, "Most Frequent Pattern": 80},
                formats={"Errors": FMT_INT, "% of Total": FMT_PCT},
                wrap=("Most Frequent Pattern",), databars=("Errors",))


def _quality_sheet(wb, a):
    ws = wb.create_sheet("Data Quality")
    ws.sheet_properties.tabColor = "7F7F7F"
    _title(ws, "Data Quality & Coverage", "How the data was collected and cleaned, and how tenants were identified.")
    s = a.stats
    total = s["rows_analyzed"] or 1
    rows = [
        ("Collection", "Query windows run", s["windows_queried"], "Each window returns at most 5,000 rows."),
        ("Collection", "Windows split for hitting the row limit", s["windows_split"],
         "A full window is halved and re-queried so no logs are dropped."),
        ("Collection", "Windows still at the row limit", s["windows_truncated"],
         "0 means the data set is complete." if not s["windows_truncated"] else "Some logs may be missing."),
        ("Cleaning", "Rows fetched from New Relic", s["rows_fetched"], ""),
        ("Cleaning", "Duplicate rows removed", s["duplicates_removed"],
         "Same timestamp, message, tenant, path and service - the same log ingested more than once."),
        ("Cleaning", "Rows analysed", s["rows_analyzed"], "Basis for every count in this workbook."),
        ("Tenant attribution", "From log context (context.TenantId)", s["tenant_from_context"],
         f"{s['tenant_from_context'] / total:.1%} of rows"),
        ("Tenant attribution", "From TenantId in message text", s["tenant_from_message"],
         f"{s['tenant_from_message'] / total:.1%} of rows"),
        ("Tenant attribution", "From identity tenant (learned)", s["tenant_from_identity"],
         f"{s['identity_mappings_learned']} identity tenant IDs matched to CoreRelate tenants using rows that carry both IDs."),
        ("Tenant attribution", "From user e-mail domain", s["tenant_from_domain"],
         "Domain learned from the data or listed in DOMAIN_MAP (tenants.py)."),
        ("Tenant attribution", "Identity tenant not resolved", s["tenant_unresolved_identity"],
         "Listed below - add the domain to DOMAIN_MAP to resolve."),
        ("Tenant attribution", "No tenant information", s["tenant_unattributed"],
         f"{s['tenant_unattributed'] / total:.1%} of rows - e.g. anonymous requests and HTTP 500 lines."),
        ("Tenant attribution", "Tenant IDs missing from TENANT_MAP", len(s["unmapped_tenant_ids"]),
         ", ".join(s["unmapped_tenant_ids"]) or "None - every tenant ID found has a name."),
    ]
    df = pd.DataFrame(rows, columns=["Area", "Check", "Value", "Notes"])
    top, last = write_table(ws, df, "DataQuality", widths={"Area": 40, "Check": 42, "Value": 12, "Notes": 90},
                            formats={"Value": FMT_INT}, wrap=("Notes",), freeze=False)

    r = last + 3
    ws.cell(r, 1, "Unresolved Identity Tenants").font = SECTION_FONT
    ws.cell(r + 1, 1, "Identity-provider tenant IDs seen in 'User: ... Tenant: ...' lines that could not be "
                      "linked to a CoreRelate tenant.").font = SUBTITLE_FONT
    un = a.unresolved_identities.rename(columns={"IdentityTenantId": "Identity Tenant ID",
                                                 "UserDomains": "User E-mail Domains", "LastSeen": "Last Seen (UTC)"})
    if un.empty:
        un = pd.DataFrame([["(none)", "", 0, None]], columns=["Identity Tenant ID", "User E-mail Domains",
                                                             "Errors", "Last Seen (UTC)"])
    un = un[["Identity Tenant ID", "User E-mail Domains", "Errors", "Last Seen (UTC)"]]
    t2, l2 = write_table(ws, un, "UnresolvedIdentities", top=r + 3, freeze=False, set_widths=False,
                         formats={"Errors": FMT_INT, "Last Seen (UTC)": FMT_DT})
    for rr in range(t2 + 1, l2 + 1):
        ws.cell(rr, 4).alignment = Alignment(horizontal="left", vertical="top")
    ws.freeze_panes = None


SHEET_GUIDE = [
    ("Overview", "Dashboard: headline numbers, key findings, errors by category and a daily trend.",
     "Start here. Every number is after duplicates are removed."),
    ("Charts", "Hourly trend and top-15 charts for patterns, tenants and endpoints.",
     "Spot spikes and the biggest contributors at a glance."),
    ("Error Patterns", "One row per distinct error type with counts, affected tenants and a sample message.",
     "Replaces the old Error Summary. Filter by Category; use Pattern ID to find rows in Raw Logs."),
    ("Tenants", "Errors per tenant with its most frequent pattern and endpoint.",
     "See which customers are affected most."),
    ("Tenant Top Errors", "The five most frequent patterns for each tenant.",
     "Filter the Tenant column to review one customer."),
    ("Tenant Daily", "Tenant x day matrix, shaded by volume.",
     "Darker cells = more errors. Shows whether a tenant's issue is ongoing or a one-day spike."),
    ("Endpoints", "Errors per API route with IDs collapsed to {id}.",
     "Find the routes that fail most."),
    ("Daily Trend", "Errors per day split by category.", "Data behind the Overview daily chart."),
    ("Hourly Heatmap", "Day x hour-of-day matrix (UTC), shaded by volume.",
     "Find the time of day errors cluster."),
    ("Data Quality", "Collection, de-duplication and tenant attribution statistics.",
     "Check completeness before drawing conclusions."),
    ("Raw Logs", "Every de-duplicated log row, newest first, with category, pattern and tenant added.",
     "Filter by Pattern ID, Tenant or Category to see the underlying events."),
]


def _guide_sheet(wb):
    ws = wb.create_sheet("Guide")
    ws.sheet_properties.tabColor = "7F7F7F"
    _title(ws, "How to Use This Workbook", "What each sheet contains. Categories and their meaning are listed below.")
    df = pd.DataFrame(SHEET_GUIDE, columns=["Sheet", "What it shows", "How to use it"])
    top, last = write_table(ws, df, "Guide", widths={"Sheet": 20, "What it shows": 70, "How to use it": 70},
                            wrap=("What it shows", "How to use it"), freeze=False)
    cats = pd.DataFrame(list(CATEGORY_DESCRIPTIONS.items()), columns=["Category", "Meaning"])
    r = last + 3
    ws.cell(r, 1, "Error Categories").font = SECTION_FONT
    t2, l2 = write_table(ws, cats, "CategoryGuide", top=r + 1, wrap=("Meaning",), freeze=False, set_widths=False)
    for i, cat in enumerate(cats["Category"]):
        ws.cell(t2 + 1 + i, 1).border = Border(left=Side(style="thick", color=CATEGORY_COLORS.get(cat, ACCENT)))


def _raw_sheet(wb, a):
    ws = wb.create_sheet("Raw Logs")
    ws.sheet_properties.tabColor = "7F7F7F"
    _title(ws, "Raw Logs", f"{len(a.raw):,} de-duplicated ERROR log rows, newest first. Timestamps are UTC.")
    write_table(ws, a.raw, "RawLogs",
                widths={"Timestamp (UTC)": 20, "Tenant": 30, "Tenant ID": 38, "Tenant Source": 22,
                        "Category": 20, "Pattern ID": 10, "Method": 8, "Status": 8, "Request Path": 45,
                        "Service": 16, "Level": 8, "Message": 100},
                formats={"Timestamp (UTC)": "yyyy-mm-dd hh:mm:ss.000"})


def build_excel(a):
    wb = Workbook()
    wb.active.title = "Overview"

    # Detail sheets first - the Overview and Charts sheets chart their data
    _patterns_sheet(wb, a)
    _tenants_sheet(wb, a)
    _tenant_top_sheet(wb, a)
    _matrix_sheet(wb, a.tenant_daily.rename(columns={"TenantName": "Tenant"}), "Tenant Daily", "TenantDaily",
                  "Tenant Errors by Day", "Errors per tenant per day (UTC). Darker = more errors.", 34, "1BAF7A")
    _endpoints_sheet(wb, a)
    _matrix_sheet(wb, a.daily, "Daily Trend", "DailyTrend", "Daily Trend by Category",
                  "Errors per day (UTC) split by category.", 18, "EDA100")
    _matrix_sheet(wb, a.heatmap, "Hourly Heatmap", "HourlyHeatmap", "Hourly Heatmap",
                  "Errors by day and hour of day (UTC). Darker = more errors.", 18, "EDA100")
    _quality_sheet(wb, a)
    _guide_sheet(wb)
    _raw_sheet(wb, a)

    data = _chart_data(wb, a)
    _overview(wb, a, data)
    _charts_sheet(wb, a, data)

    out = BytesIO()
    wb.save(out)
    out.seek(0)
    return out
